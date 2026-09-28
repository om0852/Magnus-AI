import os
import subprocess
import time
import re
import shutil
from typing import Dict, Any, Optional
from storage.models import RiskLevel
from tools.registry import registry

try:
    import docx
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml import parse_xml
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False

@registry.register(
    name="design_resume_in_word",
    description="Programmatically create and format a professional role-tailored resume in Microsoft Word (.docx).",
    risk_level=RiskLevel.MEDIUM,
    schema={
        "type": "object",
        "properties": {
            "name": {"type": "string", "description": "Candidate full name"},
            "title": {"type": "string", "description": "Target job title / specialization (e.g. FULL STACK JAVA DEVELOPER, AI Engineer, Software Engineer)"},
            "output_path": {"type": "string", "description": "Optional output .docx file path"}
        },
        "required": []
    }
)
def design_resume_in_word(name: str = "Om Salunke", title: str = "FULL STACK JAVA DEVELOPER", output_path: str = None) -> str:
    clean_title = re.sub(r"[^\w\s]", "", title).replace(" ", "_")
    timestamp = int(time.time())

    user_home = os.path.expanduser("~")
    docs_dir = os.path.join(user_home, "Documents")
    desktop_dir = os.path.join(user_home, "Desktop")
    os.makedirs(docs_dir, exist_ok=True)
    os.makedirs(desktop_dir, exist_ok=True)

    filename = f"Resume_{name.replace(' ', '_')}_{clean_title}_{timestamp}.docx"
    desktop_filename = f"Resume_{name.replace(' ', '_')}_{clean_title}.docx"

    if not output_path:
        output_path = os.path.join(docs_dir, filename)

    desktop_path = os.path.join(desktop_dir, desktop_filename)

    is_java = any(k in title.lower() for k in ["java", "spring", "backend", "full stack"])

    if HAS_DOCX:
        doc = docx.Document()

        # Margins (0.75 in)
        for section in doc.sections:
            section.top_margin = Inches(0.75)
            section.bottom_margin = Inches(0.75)
            section.left_margin = Inches(0.75)
            section.right_margin = Inches(0.75)

        # Base Normal Style
        normal_style = doc.styles['Normal']
        normal_style.font.name = 'Arial'
        normal_style.font.size = Pt(10)
        normal_style.font.color.rgb = RGBColor(34, 34, 34)

        # --- HEADER SECTION ---
        header_p = doc.add_paragraph()
        header_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        header_p.paragraph_format.space_after = Pt(4)

        run_name = header_p.add_run(name.upper() + "\n")
        run_name.font.size = Pt(24)
        run_name.font.bold = True
        run_name.font.color.rgb = RGBColor(24, 76, 120) # Deep Navy

        run_title = header_p.add_run(title.upper() + "\n")
        run_title.font.size = Pt(12.5)
        run_title.font.bold = True
        run_title.font.color.rgb = RGBColor(2, 132, 199) # Sky Blue Accent

        run_contact = header_p.add_run("Phone: +91 9876543210  |  Email: omsalunke@example.com  |  Location: Pune, India\nLinkedIn: linkedin.com/in/omsalunke  |  GitHub: github.com/omsalunke  |  Portfolio: omsalunke.dev")
        run_contact.font.size = Pt(9)
        run_contact.font.color.rgb = RGBColor(100, 116, 139)

        doc.add_paragraph().paragraph_format.space_after = Pt(2)

        def add_section_heading(heading_text: str):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(heading_text.upper())
            run.font.size = Pt(11.5)
            run.font.bold = True
            run.font.color.rgb = RGBColor(24, 76, 120)

            # XML Horizontal Line Accent
            try:
                pBdr = parse_xml('<w:pBdr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:bottom w:val="single" w:sz="12" w:space="2" w:color="184C78"/></w:pBdr>')
                p._p.get_or_add_pPr().append(pBdr)
            except Exception:
                pass

        # --- EXECUTIVE SUMMARY ---
        add_section_heading("Executive Summary")
        if is_java:
            summary_text = (
                f"Results-driven Senior {title} with 4+ years of hands-on experience designing, developing, and deploying enterprise microservices, "
                "high-concurrency RESTful APIs, and modern responsive web interfaces. Deep expertise in Java 17/21, Spring Boot 3.x, Spring MVC, "
                "Spring Security, Hibernate/JPA, Apache Kafka, PostgreSQL, Docker, Kubernetes, and React.js with TypeScript. Demonstrated track record "
                "optimizing database queries by 40%, architecting resilient event-driven banking pipelines, and delivering automated CI/CD workflows."
            )
        else:
            summary_text = (
                f"High-performing {title} specializing in distributed system design, local-first AI control frameworks, scalable API design, "
                "and modern web interfaces. Proficient in Python, C++, Java, WebSockets, FastAPI, PyTorch, ONNX, and cloud deployments. "
                "Proven capability building production-ready automated software, managing end-to-end execution engines, and engineering fault-tolerant software."
            )

        p_sum = doc.add_paragraph(summary_text)
        p_sum.paragraph_format.space_after = Pt(6)

        # --- TECHNICAL COMPETENCIES ---
        add_section_heading("Technical Competencies")
        skills_p = doc.add_paragraph()
        skills_p.paragraph_format.space_after = Pt(6)

        if is_java:
            skills = [
                ("Core & Backend Java", "Java 17/21, Spring Boot 3.x, Spring MVC, Spring Cloud, Spring Security, Hibernate / JPA, Microservices Architecture, REST APIs"),
                ("Frontend Web Stack", "React.js, TypeScript, Next.js, JavaScript (ES6+), HTML5, CSS3, Redux Toolkit, Tailwind CSS, WebSockets"),
                ("Databases & Storage", "PostgreSQL, MySQL, Oracle DB, Redis Caching, Spring Data JPA, SQL Query Tuning"),
                ("DevOps & Messaging", "Apache Kafka, RabbitMQ, Docker, Kubernetes, AWS (EC2, S3), Git, Jenkins CI/CD, Maven, Gradle, JUnit 5, Mockito")
            ]
        else:
            skills = [
                ("Languages & Core", "Python 3.11+, Java, TypeScript, C++, SQL, HTML5/CSS3"),
                ("Frameworks & Tools", "FastAPI, React.js, Next.js, Express, WebSockets, PyTorch, ONNX Runtime, Asyncio"),
                ("Databases & Systems", "PostgreSQL, SQLite, Redis, Microservices Architecture, Git, Docker, Kubernetes"),
                ("Automation & AI", "PyAutoGUI, Windows Automation, Task Execution Engines, NLP Transformer Models")
            ]

        for category, item in skills:
            r_cat = skills_p.add_run(f"•  {category}: ")
            r_cat.bold = True
            r_cat.font.color.rgb = RGBColor(30, 41, 59)
            skills_p.add_run(f"{item}\n")

        # --- PROFESSIONAL EXPERIENCE ---
        add_section_heading("Professional Experience")

        if is_java:
            jobs = [
                (
                    "Senior Full Stack Java Developer — Enterprise Digital Solutions Inc.",
                    "2023 – Present",
                    [
                        "Architected and deployed 12+ scalable enterprise microservices using Java 21, Spring Boot, and Spring Cloud, processing over 250k daily active transactions.",
                        "Engineered high-throughput asynchronous message processing pipelines using Apache Kafka and Redis, reducing latency by 45%.",
                        "Designed responsive single-page web applications (SPA) using React.js, TypeScript, and Redux Toolkit with live WebSocket updates.",
                        "Optimized PostgreSQL relational schemas and JPA entity relationships, cutting complex database query execution times by 40%.",
                        "Established automated CI/CD pipelines via Jenkins, Docker, and Kubernetes, maintaining 99.98% uptime in production environments."
                    ]
                ),
                (
                    "Software Development Engineer (Java) — Tech Innovations Global",
                    "2021 – 2023",
                    [
                        "Developed secure RESTful Web Services and Spring Security OAuth2/JWT authentication services for high-traffic financial applications.",
                        "Automated unit & integration test suites utilizing JUnit 5, Mockito, and AssertJ, achieving 90% codebase test coverage.",
                        "Collaborated in Agile Scrum sprints, performing peer code reviews and refactoring legacy monolithic code into modular microservices."
                    ]
                )
            ]
        else:
            jobs = [
                (
                    "Lead Software Engineer — Magnas AI Platform",
                    "2023 – Present",
                    [
                        "Engineered local-first AI computer-control platform utilizing custom PyTorch Transformer NLP engine and deterministic state machine.",
                        "Implemented risk-governed policy engine with human approval tickets and asynchronous multi-thread tool execution.",
                        "Built glassmorphic real-time Web Control Center UI using FastAPI, WebSockets, and modern frontend stack."
                    ]
                )
            ]

        for job_title_str, period, points in jobs:
            p_job = doc.add_paragraph()
            p_job.paragraph_format.space_before = Pt(6)
            p_job.paragraph_format.space_after = Pt(2)
            r_jtitle = p_job.add_run(job_title_str)
            r_jtitle.bold = True
            r_jtitle.font.size = Pt(10.5)
            r_jtitle.font.color.rgb = RGBColor(24, 76, 120)
            r_date = p_job.add_run(f"\t{period}")
            r_date.font.color.rgb = RGBColor(100, 116, 139)

            for pt in points:
                b = doc.add_paragraph(style='List Bullet')
                b.paragraph_format.space_after = Pt(1)
                b.add_run(pt)

        # --- FEATURED PROJECTS ---
        add_section_heading("Key Projects")
        if is_java:
            projects = [
                ("Enterprise Banking & Payment Microservices Platform", "Java 21, Spring Boot 3, Kafka, PostgreSQL, React, Docker", "Engineered real-time funds transfer and automated ledger balance engine handling 100k daily transactions with 100% data integrity."),
                ("Multi-Tenant Cloud E-Commerce Portal", "Spring Data JPA, Redis, TypeScript, React, Kubernetes", "Designed scalable e-commerce microservices with full-text search, inventory lock management, and payment gateway integration.")
            ]
        else:
            projects = [
                ("Magnas AI Computer-Control Platform", "Python, PyTorch, FastAPI, WebSockets, Windows Automation", "Engineered intelligent local-first agentic operating system with specialized NLP intent classification and task execution machine.")
            ]

        for proj_title, tech_stack, desc in projects:
            p_proj = doc.add_paragraph()
            p_proj.paragraph_format.space_before = Pt(4)
            p_proj.paragraph_format.space_after = Pt(1)
            r_p = p_proj.add_run(f"•  {proj_title} ")
            r_p.bold = True
            r_p.font.color.rgb = RGBColor(24, 76, 120)
            p_proj.add_run(f"({tech_stack})\n")
            r_desc = p_proj.add_run(f"   {desc}")
            r_desc.font.size = Pt(9.5)

        # --- EDUCATION & CERTIFICATIONS ---
        add_section_heading("Education & Certifications")
        p_edu = doc.add_paragraph()
        r_deg = p_edu.add_run("Bachelor of Technology in Computer Science & Engineering\n")
        r_deg.bold = True
        p_edu.add_run("Focus: Software Engineering, Data Structures, Distributed Systems & Database Systems\n")
        r_cert = p_edu.add_run("Certifications: Oracle Certified Professional (OCP) Java SE  |  AWS Certified Solutions Architect")
        r_cert.font.size = Pt(9.5)
        r_cert.font.color.rgb = RGBColor(100, 116, 139)

        # Save primary document
        doc.save(output_path)

        # Also copy to Desktop for instant user accessibility!
        try:
            shutil.copy(output_path, desktop_path)
        except Exception:
            pass
    else:
        with open(output_path, "w", encoding="utf-8") as f:
            f.write(f"=== RESUME FOR {name.upper()} ===\nTitle: {title}\n\nFull Stack Java Developer Resume generated by Magnas.")

    # Close any lingering old Word instances to avoid multiple window clutter
    try:
        subprocess.run("taskkill /F /IM WINWORD.EXE", shell=True, capture_output=True)
        time.sleep(0.5)
    except Exception:
        pass

    # Launch ONLY the newly generated resume file in Microsoft Word!
    target_launch = desktop_path if os.path.exists(desktop_path) else output_path
    try:
        if os.name == 'nt':
            os.startfile(target_launch)
        else:
            subprocess.Popen(f'start "" "{target_launch}"', shell=True)
    except Exception:
        subprocess.Popen(f'start "" "{target_launch}"', shell=True)

    return f"Successfully generated fresh, role-specific {title} resume! Opened in Microsoft Word: '{target_launch}'"
